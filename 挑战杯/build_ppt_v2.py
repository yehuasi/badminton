# -*- coding: utf-8 -*-
"""第二版答辩演示：直接在「演示1.pptx」原页面上改文字，保留模板版式设计。

做法：封面/目录/章节页/致谢只替换文字；正文页保留顶部「项目重点」标题栏，
删除正文区的信息图内容，再填入本项目内容与配图。
输出：挑战杯/答辩演示_v2.pptx
"""

import os

import build_ppt_from_template as b
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "演示1.pptx")
OUT = os.path.join(HERE, "答辩演示_v2.pptx")

ACCENT = RGBColor(0x07, 0x65, 0xED)


def replace_full(slide, old_sub, new_full):
    """找到包含 old_sub 的文本框，替换为 new_full（保留首段样式）。"""
    for sh in slide.shapes:
        if sh.has_text_frame and old_sub in sh.text_frame.text:
            b.set_ph_text(sh, new_full)
            return True
    return False


def transform_content(slide, title, cfg):
    """正文页：保留顶部标题栏，删除正文信息图，填入本项目内容。"""
    # 删除正文区（y >= 1.15 英寸）
    for sh in list(slide.shapes):
        if sh.top >= Inches(1.15):
            sh._element.getparent().remove(sh._element)
    # 标题栏里替换标题、清空英文副标题
    for sh in list(slide.shapes):
        if sh.top >= Inches(1.15) or not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if not t or t == "项目重点":
            continue
        if any("\u4e00" <= c <= "\u9fff" for c in t):
            b.set_ph_text(sh, title)
        elif t.isascii():
            b.set_ph_text(sh, "")

    if "images" in cfg:
        b.add_image_grid(slide, cfg["images"])
        return
    has_fig = cfg.get("figure") is not None
    if cfg.get("photo"):
        b.add_textbox(slide, 0.7, 1.4, 7.2, 5.4, cfg["bullets"], size=15)
        b.add_picture_fit(slide, cfg["photo"], 8.15, 1.45, 4.7, 5.5)
        return
    b.add_textbox(slide, 0.8, 1.3, 11.7, 2.7 if has_fig else 5.5, cfg["bullets"], size=16)
    if has_fig:
        b.add_figure(slide, cfg["figure"])


def main():
    prs = Presentation(TEMPLATE)
    slides = list(prs.slides)

    # ---- 封面 ----
    s = slides[0]
    replace_full(s, "AI数据建模场景", "轻羽智评")
    replace_full(s, "大学生创新创业计划书", "端侧人工智能羽毛球动作识别与规范评估系统")
    replace_full(s, "College Student Innovation", "Smart Badminton · Edge AI Action Assessment")
    replace_full(s, "20XX", "汇报时间：2026.10")
    replace_full(s, "OfficePLUS", "汇报人：XXXX、XXXX")

    # ---- 目录 ----
    s = slides[1]
    for old, new in [
        ("建模场景需求分析", "项目背景与总体方案"),
        ("Modeling scenario requirements analysis", "Background and overall design"),
        ("AI模型选择与优化", "核心算法与实验数据"),
        ("AI model selection and optimization", "Core algorithms and experiments"),
        ("数据安全与隐私保护", "端侧部署与系统演示"),
        ("Data security and privacy protection", "Edge deployment and demo"),
        ("实施与部署策略", "应用推广与未来规划"),
        ("Implementation and deployment strategies", "Applications and roadmap"),
    ]:
        replace_full(s, old, new)

    # ---- 章节页 ----
    for idx, old_zh, old_en, new_zh, new_en, new_desc in [
        (2, "建模场景需求分析", "Modeling scenario requirements analysis",
         "项目背景与总体方案", "Background and overall design",
         "面向普通人动作难以自查的痛点，给出端侧 AI 动作识别与规范评估的整体方案。"),
        (6, "AI模型选择与优化", "AI model selection and optimization",
         "核心算法与实验数据", "Core algorithms and experiments",
         "七类动作识别、分组门控、四项可解释评分，以及目标域微调后的实测结果。"),
        (10, "数据安全与隐私保护", "Data security and privacy protection",
         "端侧部署与系统演示", "Edge deployment and demo",
         "SC171 端侧部署、四项 Web 视图与性能对比，支持现场实物演示。"),
        (14, "实施与部署策略", "Implementation and deployment strategies",
         "应用推广与未来规划", "Applications and roadmap",
         "应用场景、低成本平台迁移路径与后续技术演进。"),
    ]:
        s = slides[idx]
        replace_full(s, old_zh, new_zh)
        replace_full(s, old_en, new_en)
        replace_full(s, "利用A构建的信用评分模型", new_desc)

    # ---- 正文 ----
    content = [
        dict(title="项目背景与痛点",
             bullets=["羽毛球普及度高，动作规范却难以自查",
                      "专业教练稀缺、费用较高，普通人缺少及时反馈",
                      "市面方案多依赖云端，存在隐私与网络依赖",
                      "目标：单摄像头加端侧 AI，做随时可用的动作教练"]),
        dict(title="现有技术对比",
             bullets=["人工录像回看：无量化指标，依赖个人经验",
                      "云端动作分析：需联网，隐私风险，评分不透明",
                      "本系统：端侧本地运行，实时量化，建议可解释"]),
        dict(title="系统总体设计",
             bullets=["三级流水线：视频流粗筛 → 姿态关键点 → 识别与评分",
                      "一套算法、两种入口：实时与录像共用",
                      "局域网 Web：实时、回看、报告、双视角融合"],
             figure="fig_pipeline.png"),
        dict(title="核心技术与创新点",
             bullets=["活动触发式切分：一拍一次，快慢挥拍自适应",
                      "识别门控：无动作不强行归类",
                      "AHP 可解释评分：权重可追溯，建议可解释",
                      "目标域微调：同人同机位识别 87.9%",
                      "双视角融合：正面加侧面分时采集"]),
        dict(title="动作识别与实验数据",
             bullets=["七类击球：正反手发球、高远球、杀球、吊球、挑球、搓球",
                      "过顶/非过顶分组门控，抑制跨组误判",
                      "域适配后同人同机位留出准确率 87.9% ± 3.6%"],
             figure="fig_gate.png"),
        dict(title="四项评分与建议",
             bullets=["动作标准度 50%、步伐 25%、连贯性 15%、动作识别 10%",
                      "识别置信度低于 0.5 时动作分减半",
                      "逐条错误规则加纠正话术，评分可解释"],
             figure="fig_score.png"),
        dict(title="端侧部署与性能",
             bullets=["SC171 v3（QCM6490），ONNX Runtime 纯 CPU",
                      "显示 30fps，分析帧率 13~17fps",
                      "分类器 ONNX 约 0.36MB，姿态模型约 13MB",
                      "切分命中率 95.8%"],
             photo="photo/开发板实物照片.jpg"),
        dict(title="系统界面",
             images=[("photo/实时.png", "实时识别与建议"),
                     ("photo/回看.png", "逐帧回看"),
                     ("photo/报告.png", "会话报告"),
                     ("photo/融合.png", "双视角融合")]),
        dict(title="性能对比与优势验证",
             bullets=["端侧本地运行，无需联网，保护隐私",
                      "量化评分加可解释建议，区别于黑盒打分",
                      "实物可在现场演示，非纸面方案"],
             figure="fig_perf.png"),
        dict(title="应用场景与推广价值",
             bullets=["体校与俱乐部教学辅助、中小学体育课、居家训练、康复",
                      "单摄像头加开发板，硬件成本低",
                      "平台解耦：SC171 验证可行性，可迁移到 RK3588 等更低成本平台",
                      "同一套框架可迁移到健身、网球、乒乓球"]),
        dict(title="应用推广与未来规划",
             bullets=["已完成数据、训练、评分、部署全链路",
                      "未来：NPU 加速、更多运动、专业示范数据集",
                      "社会价值：普及科学健身、缓解教练不足"]),
    ]
    slot = [3, 4, 5, 7, 8, 9, 11, 12, 13, 15, 16]
    for cfg, idx in zip(content, slot):
        transform_content(slides[idx], cfg["title"], cfg)

    # 删除多余的一个正文页（0-based 17）
    xml_slides = prs.slides._sldIdLst
    all_slds = list(xml_slides)
    dead = all_slds[17]
    prs.part.drop_rel(dead.get(qn("r:id")))
    xml_slides.remove(dead)

    # ---- 致谢 ----
    s = slides[18]
    replace_full(s, "AI数据建模场景", "轻羽智评")
    replace_full(s, "College Student Innovation", "Smart Badminton · Edge AI Action Assessment")
    replace_full(s, "OfficePLUS", "汇报人：XXXX、XXXX")
    replace_full(s, "20XX", "汇报时间：2026.10")

    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
